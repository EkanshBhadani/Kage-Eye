# report_writer.py

import os

from openpyxl import Workbook
from openpyxl import load_workbook


class ReportWriter:

    def __init__(
        self,
        filepath="data/results.xlsx"
    ):
        self.filepath = filepath

    def create_fresh_report(self):

        os.makedirs(
            os.path.dirname(self.filepath),
            exist_ok=True
        )

        # Delete old report
        if os.path.exists(self.filepath):
            os.remove(self.filepath)

        wb = Workbook()

        ws = wb.active
        ws.title = "Results"

        # ======================
        # PASSED SECTION
        # ======================

        ws["C1"] = "PASSED"

        ws["A2"] = "Username"
        ws["B2"] = "Score"
        ws["C2"] = "Time stamp (DD-MM-YYYY HH:MM:SS)"
        ws["D2"] = "Reply No."

        # ======================
        # FAILED SECTION
        # ======================

        ws["H1"] = "FAILED"

        ws["F2"] = "Username"
        ws["G2"] = "Score"
        ws["H2"] = "Time stamp (DD-MM-YYYY HH:MM:SS)"
        ws["I2"] = "Reply No."

        # ======================
        # MISSING SECTION
        # ======================

        ws["M1"] = "MISSING"

        ws["K2"] = "Username"
        ws["L2"] = "Score"
        ws["M2"] = "Time stamp (DD-MM-YYYY HH:MM:SS)"
        ws["N2"] = "Reply No."

        # ======================
        # SUMMARY SHEET
        # ======================

        summary = wb.create_sheet("Summary")

        summary["A1"] = "Metric"
        summary["B1"] = "Value"

        # ======================
        # AUDIT LOG SHEET
        # ======================

        audit = wb.create_sheet("Audit Log")

        audit.append([
            "Reply Number",
            "Forum User",
            "Submitted Username",
            "Score",
            "Result",
            "Form Timestamp"
        ])

        wb.save(self.filepath)

    def _next_passed_row(self, ws):

        row = 3

        while ws[f"A{row}"].value:
            row += 1

        return row

    def _next_failed_row(self, ws):

        row = 3

        while ws[f"F{row}"].value:
            row += 1

        return row

    def _next_missing_row(self, ws):

        row = 3

        while ws[f"K{row}"].value:
            row += 1

        return row

    def add_passed(
        self,
        username,
        score,
        timestamp,
        reply_number
    ):

        wb = load_workbook(self.filepath)

        ws = wb["Results"]

        row = self._next_passed_row(ws)

        ws[f"A{row}"] = f"@{username}"
        ws[f"B{row}"] = score
        ws[f"C{row}"] = timestamp
        ws[f"D{row}"] = reply_number

        wb.save(self.filepath)

    def add_failed(
        self,
        username,
        score,
        timestamp,
        reply_number
    ):

        wb = load_workbook(self.filepath)

        ws = wb["Results"]

        row = self._next_failed_row(ws)

        ws[f"F{row}"] = f"@{username}"
        ws[f"G{row}"] = score
        ws[f"H{row}"] = timestamp
        ws[f"I{row}"] = reply_number

        wb.save(self.filepath)

    def add_missing(
        self,
        username,
        reply_number
    ):

        wb = load_workbook(self.filepath)

        ws = wb["Results"]

        row = self._next_missing_row(ws)

        ws[f"K{row}"] = f"@{username}"
        ws[f"L{row}"] = "N/A"
        ws[f"M{row}"] = "N/A"
        ws[f"N{row}"] = reply_number

        wb.save(self.filepath)

    def add_audit(
        self,
        reply_number,
        forum_user,
        username,
        score,
        result,
        timestamp
    ):

        wb = load_workbook(self.filepath)

        audit = wb["Audit Log"]

        audit.append([
            reply_number,
            forum_user,
            username,
            score,
            result,
            timestamp
        ])

        wb.save(self.filepath)

    def update_summary(self):

        wb = load_workbook(self.filepath)

        results = wb["Results"]
        summary = wb["Summary"]

        passed = 0
        failed = 0
        missing = 0

        row = 3

        while results[f"A{row}"].value:
            passed += 1
            row += 1

        row = 3

        while results[f"F{row}"].value:
            failed += 1
            row += 1

        row = 3

        while results[f"K{row}"].value:
            missing += 1
            row += 1

        if summary.max_row > 1:
            summary.delete_rows(
                2,
                summary.max_row
            )

        summary.append([
            "Passed",
            passed
        ])

        summary.append([
            "Failed",
            failed
        ])

        summary.append([
            "Missing",
            missing
        ])

        summary.append([
            "Total Processed",
            passed + failed + missing
        ])

        wb.save(self.filepath)
